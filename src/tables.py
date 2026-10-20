"""
Table S1 (all included publications, with standardised fields) and the
mapping QA workbook.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.data_loading import OPTIONAL_COLUMNS

_join = '; '.join


def _metric(row) -> str:
    if pd.isna(row['metric_type']):
        return ''
    lo, hi = row['metric_min'], row['metric_max']
    if pd.isna(hi):
        return row['metric_type']
    if pd.isna(lo) or lo == hi:
        return f"{row['metric_type']} {hi:g}"
    return f"{row['metric_type']} {lo:g}–{hi:g}"


def table_s1(df: pd.DataFrame) -> pd.DataFrame:
    table = pd.DataFrame({
        'ID': df['ID'],
        'Authors': df['Authors'],
        'Year': df['Year'],
        'Title': df['Title'],
        'DOI': df['DOI'],
        'Journal': df['Journal'],
        'Review': df['is_review'].map({True: 'yes', False: ''}),
        'WOAH-listed viral diseases': df['disease_list'].map(_join),
        'Diseases (as extracted)': df['Diseases'],
        'Algorithm categories (F1)': df['F1'].map(_join),
        'Algorithms (F2, standardised)': df['F2'].map(_join),
        'Countries': df['Countries'],
        'Stage': df['Stage'],
        'Theme': df['Theme'],
        'Target': df['Target'],
        'Main metric': df.apply(_metric, axis=1),
    })
    for col in OPTIONAL_COLUMNS:
        if col in df.columns:
            table[col] = df[col]
    table['Notes'] = df['Notes']
    table['Sources'] = df['Sources']
    return table.sort_values(['Year', 'Authors']).reset_index(drop=True)


COLUMN_WIDTHS = {
    'ID': 6, 'Authors': 30, 'Year': 7, 'Title': 50, 'DOI': 26, 'Journal': 28,
    'Review': 8, 'WOAH-listed viral diseases': 18, 'Diseases (as extracted)': 24,
    'Algorithm categories (F1)': 26, 'Algorithms (F2, standardised)': 40,
    'Countries': 22, 'Stage': 16, 'Theme': 24, 'Target': 24, 'Main metric': 18,
    'Notes': 60, 'Sources': 16,
}


def save_table_s1(df: pd.DataFrame, output_path: Path) -> None:
    table = table_s1(df)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Table S1'

    header_fill = PatternFill(start_color='2E4057', end_color='2E4057', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    border = Border(*(Side(style='thin'),) * 4)

    for col, name in enumerate(table.columns, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill, cell.font, cell.border = header_fill, header_font, border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS.get(name, 14)
    for r, row in enumerate(table.itertuples(index=False), 2):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=col, value=None if pd.isna(value) else value)
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f'  Saved: {output_path} ({len(table)} publications)')


def save_qa_report(sheets: dict, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        for name, sheet in sheets.items():
            sheet.to_excel(writer, sheet_name=name[:31], index=False)
    print(f'  Saved: {output_path}')

"""
Table S2 (complete bibliography) generation for PRISMA-ScR scoping review.
One row per publication (no algorithm/disease/country expansion).
"""

import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from src.standardization import apply_publication_level_standardization

COLUMNS_ORDER = [
    'Authors', 'Year', 'Title', 'DOI', 'Journal',
    'Disease', 'Disease_standardized',
    'Algorithm(s)', 'Algorithm_standardized',
    'Section',
    'Country', 'Country_standardized',
    'Main metric',
    'Stage', 'Stage_standardized',
    'Sources', 'Notes',
]

COLUMN_WIDTHS = {
    'Authors': 16, 'Year': 6, 'Title': 45, 'DOI': 18, 'Journal': 20,
    'Disease': 14, 'Disease_standardized': 16,
    'Algorithm(s)': 16, 'Algorithm_standardized': 16,
    'Section': 18,
    'Country': 12, 'Country_standardized': 14,
    'Main metric': 12,
    'Stage': 12, 'Stage_standardized': 16,
    'Sources': 10, 'Notes': 25,
}


def generate_table_s2(df: pd.DataFrame, output_path: Path) -> None:
    """
    Build Table S2: one row per publication with standardized columns added.

    Args:
        df: Included publications (not expanded by algorithm/disease/country).
        output_path: Destination .xlsx path.
    """
    df_final = apply_publication_level_standardization(df)

    available_cols = [col for col in COLUMNS_ORDER if col in df_final.columns]
    df_final = df_final[available_cols].reset_index(drop=True)
    df_final.index = df_final.index + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Table S2"

    header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    for col_num, col_name in enumerate(available_cols, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = col_name
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    for row_idx, row in df_final.iterrows():
        for col_num, col_name in enumerate(available_cols, 1):
            cell = ws.cell(row=row_idx + 1, column=col_num)
            cell.value = row[col_name]
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    for col_num, col_name in enumerate(available_cols, 1):
        width = COLUMN_WIDTHS.get(col_name, 12)
        ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = width

    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"  Publications in table: {len(df_final)}")
    print(f"  Columns: {len(available_cols)}")


def save_standardized_preview(sheets: dict, output_path: Path) -> None:
    """
    Dump the standardized DataFrames actually used to draw the figures into
    one workbook (one sheet per DataFrame), for manual spot-checking of the
    disease/algorithm/country/stage mappings.

    Args:
        sheets: {sheet_name: DataFrame}, e.g. the outputs of
            apply_standardization, apply_algorithm_standardization,
            apply_algorithm_disease_standardization, and
            apply_country_standardization.
        output_path: Destination .xlsx path.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet_name[:31], index=False)

    print(f"  Sheets: {', '.join(sheets.keys())}")
